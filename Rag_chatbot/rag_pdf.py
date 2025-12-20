#!/usr/bin/env python3
"""
Standalone RAG (Retrieval-Augmented Generation) Knowledge Base System
No Streamlit interface - Direct code-based interaction with knowledge base folder
"""

import os
import json
import hashlib
import time
import glob
from datetime import datetime
from typing import List, Dict, Any, Optional
import sqlite3
import logging
from pathlib import Path

# Core RAG imports
from langchain_community.llms import Ollama
from langchain_community.document_loaders import PyPDFLoader, TextLoader, CSVLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document

# PDF and Excel table extraction
import pdfplumber
import pandas as pd
import openpyxl
from io import StringIO

# 🔥 Configuration
ANSWER_MODEL = "phi3"
EMBED_MODEL = "bge-m3"
KNOWLEDGE_BASE_DIR = "documents"
CACHE_DB = "rag_cache.db"
LOG_FILE = "rag_debug.log"

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler(LOG_FILE),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

class EnhancedPDFLoader:
    """Enhanced PDF loader that extracts both text and tables"""
    
    def __init__(self, file_path: str):
        self.file_path = file_path
    
    def load(self) -> List[Document]:
        """Load PDF with text and table extraction"""
        documents = []
        
        try:
            # Extract text using PyPDFLoader
            pdf_loader = PyPDFLoader(self.file_path)
            text_docs = pdf_loader.load()
            documents.extend(text_docs)
            
            # Extract tables using pdfplumber
            with pdfplumber.open(self.file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, start=1):
                    tables = page.extract_tables()
                    
                    if tables:
                        for table_num, table in enumerate(tables, start=1):
                            if table and len(table) > 0:
                                # Convert table to DataFrame for better formatting
                                df = pd.DataFrame(table[1:], columns=table[0] if table[0] else None)
                                
                                # Create structured text from table
                                table_text = f"\n=== TABLE {table_num} on Page {page_num} ===\n"
                                table_text += df.to_string(index=False)
                                table_text += "\n" + "="*50 + "\n"
                                
                                # Also create searchable key-value pairs
                                for idx, row in df.iterrows():
                                    for col in df.columns:
                                        if pd.notna(row[col]):
                                            table_text += f"{col}: {row[col]}\n"
                                
                                # Add as document
                                doc = Document(
                                    page_content=table_text,
                                    metadata={
                                        'source': self.file_path,
                                        'page': page_num,
                                        'type': 'table',
                                        'table_number': table_num
                                    }
                                )
                                documents.append(doc)
                                
        except Exception as e:
            logger.error(f"Error extracting tables from PDF {self.file_path}: {e}")
            # Fallback to regular PDF loading
            pdf_loader = PyPDFLoader(self.file_path)
            documents = pdf_loader.load()
        
        return documents


class EnhancedExcelLoader:
    """Enhanced Excel loader that handles multiple sheets and all columns"""
    
    def __init__(self, file_path: str):
        self.file_path = file_path
    
    def load(self) -> List[Document]:
        """Load all sheets and columns from Excel file"""
        documents = []
        
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(self.file_path)
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                
                # Create structured content with all columns
                sheet_content = f"\n=== EXCEL SHEET: {sheet_name} ===\n"
                sheet_content += f"Columns: {', '.join(df.columns.tolist())}\n"
                sheet_content += "="*60 + "\n\n"
                
                # Add full table view
                sheet_content += df.to_string(index=False) + "\n\n"
                
                # Create searchable key-value content for each row
                for idx, row in df.iterrows():
                    row_content = f"Row {idx+1}:\n"
                    for col in df.columns:
                        if pd.notna(row[col]):
                            row_content += f"  {col}: {row[col]}\n"
                    sheet_content += row_content + "\n"
                
                # Add sheet as document
                doc = Document(
                    page_content=sheet_content,
                    metadata={
                        'source': self.file_path,
                        'sheet': sheet_name,
                        'type': 'excel',
                        'columns': df.columns.tolist(),
                        'row_count': len(df)
                    }
                )
                documents.append(doc)
                
                # Also create column-specific documents for better searching
                for col in df.columns:
                    col_content = f"\n=== Column '{col}' from Sheet '{sheet_name}' ===\n"
                    col_values = df[col].dropna().tolist()
                    col_content += f"Values: {col_values}\n"
                    col_content += f"Unique values: {df[col].nunique()}\n"
                    
                    if df[col].dtype in ['int64', 'float64']:
                        col_content += f"Min: {df[col].min()}, Max: {df[col].max()}\n"
                        col_content += f"Mean: {df[col].mean():.2f}\n"
                    
                    col_doc = Document(
                        page_content=col_content,
                        metadata={
                            'source': self.file_path,
                            'sheet': sheet_name,
                            'type': 'excel_column',
                            'column_name': col
                        }
                    )
                    documents.append(col_doc)
                    
        except Exception as e:
            logger.error(f"Error loading Excel file {self.file_path}: {e}")
        
        return documents


class KnowledgeBaseManager:
    """Manages multiple documents in the knowledge base"""
    
    def __init__(self):
        os.makedirs(KNOWLEDGE_BASE_DIR, exist_ok=True)
        self.embeddings = OllamaEmbeddings(model=EMBED_MODEL)
        self.documents = {}
        self.combined_vectordb = None
        self.init_cache_db()
        self.load_existing_documents()
        
    def init_cache_db(self):
        """Initialize SQLite database for caching answers"""
        conn = sqlite3.connect(CACHE_DB)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS cached_answers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question_hash TEXT UNIQUE,
                question TEXT,
                answer TEXT,
                context TEXT,
                timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
                source_files TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS debug_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
                operation TEXT,
                details TEXT,
                execution_time REAL
            )
        ''')
        conn.commit()
        conn.close()
        
    def get_question_hash(self, question: str) -> str:
        """Generate hash for question to use as cache key"""
        return hashlib.md5(question.lower().strip().encode()).hexdigest()
    
    def cache_answer(self, question: str, answer: str, context: str, source_files: List[str]):
        """Cache answer for future use"""
        question_hash = self.get_question_hash(question)
        conn = sqlite3.connect(CACHE_DB)
        cursor = conn.cursor()
        try:
            cursor.execute('''
                INSERT OR REPLACE INTO cached_answers 
                (question_hash, question, answer, context, source_files)
                VALUES (?, ?, ?, ?, ?)
            ''', (question_hash, question, answer, context, json.dumps(source_files)))
            conn.commit()
            logger.info(f"Cached answer for question: {question[:50]}...")
        except Exception as e:
            logger.error(f"Error caching answer: {e}")
        finally:
            conn.close()
    
    def get_cached_answer(self, question: str) -> Optional[Dict]:
        """Retrieve cached answer if available"""
        question_hash = self.get_question_hash(question)
        conn = sqlite3.connect(CACHE_DB)
        cursor = conn.cursor()
        try:
            cursor.execute('''
                SELECT question, answer, context, source_files, timestamp
                FROM cached_answers WHERE question_hash = ?
            ''', (question_hash,))
            result = cursor.fetchone()
            if result:
                return {
                    'question': result[0],
                    'answer': result[1],
                    'context': result[2],
                    'source_files': json.loads(result[3]),
                    'timestamp': result[4]
                }
        except Exception as e:
            logger.error(f"Error retrieving cached answer: {e}")
        finally:
            conn.close()
        return None
    
    def log_operation(self, operation: str, details: str, execution_time: float):
        """Log operations for debugging"""
        conn = sqlite3.connect(CACHE_DB)
        cursor = conn.cursor()
        try:
            cursor.execute('''
                INSERT INTO debug_logs (operation, details, execution_time)
                VALUES (?, ?, ?)
            ''', (operation, details, execution_time))
            conn.commit()
        except Exception as e:
            logger.error(f"Error logging operation: {e}")
        finally:
            conn.close()
    
    def add_document(self, file_path: str, doc_type: str = "pdf") -> bool:
        """Add a document to the knowledge base"""
        start_time = time.time()
        try:
            # Use enhanced loaders for PDF and Excel
            if doc_type == "pdf":
                loader = EnhancedPDFLoader(file_path)
            elif doc_type == "xlsx" or doc_type == "excel":
                loader = EnhancedExcelLoader(file_path)
            elif doc_type == "txt":
                loader = TextLoader(file_path)
            elif doc_type == "csv":
                loader = CSVLoader(file_path)
            else:
                raise ValueError(f"Unsupported document type: {doc_type}")
            
            docs = loader.load()
            splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
            chunks = splitter.split_documents(docs)
            
            # Create vector store for this document
            doc_name = os.path.basename(file_path)
            vectordb = FAISS.from_documents(chunks, self.embeddings)
            
            # Save individual document vector store
            db_path = os.path.join(KNOWLEDGE_BASE_DIR, f"vs_{doc_name}")
            vectordb.save_local(db_path)
            
            self.documents[doc_name] = {
                'path': file_path,
                'type': doc_type,
                'db_path': db_path,
                'chunks': len(chunks),
                'added_at': datetime.now().isoformat()
            }
            
            # Rebuild combined vector store
            self._rebuild_combined_vectordb()
            
            execution_time = time.time() - start_time
            self.log_operation("add_document", f"Added {doc_name} ({doc_type})", execution_time)
            logger.info(f"Successfully added document: {doc_name}")
            return True
            
        except Exception as e:
            execution_time = time.time() - start_time
            self.log_operation("add_document_error", str(e), execution_time)
            logger.error(f"Error adding document {file_path}: {e}")
            return False
    
    def _rebuild_combined_vectordb(self):
        """Rebuild combined vector database with all documents"""
        try:
            all_vectordbs = []
            for doc_name, doc_info in self.documents.items():
                if os.path.exists(doc_info['db_path']):
                    vectordb = FAISS.load_local(
                        doc_info['db_path'],
                        self.embeddings,
                        allow_dangerous_deserialization=True
                    )
                    all_vectordbs.append(vectordb)
            
            if all_vectordbs:
                self.combined_vectordb = all_vectordbs[0]
                for vectordb in all_vectordbs[1:]:
                    self.combined_vectordb.merge_from(vectordb)
                
                # Save combined vector store
                combined_path = os.path.join(KNOWLEDGE_BASE_DIR, "combined_vs")
                self.combined_vectordb.save_local(combined_path)
                logger.info("Successfully rebuilt combined vector database")
            
        except Exception as e:
            logger.error(f"Error rebuilding combined vector database: {e}")
    
    def load_existing_documents(self):
        """Automatically scan and load documents from knowledge base folder"""
        supported_extensions = ['*.pdf', '*.txt', '*.csv', '*.xlsx', '*.xls']
        
        for extension in supported_extensions:
            pattern = os.path.join(KNOWLEDGE_BASE_DIR, extension)
            files = glob.glob(pattern)
            
            for file_path in files:
                doc_name = os.path.basename(file_path)
                db_path = os.path.join(KNOWLEDGE_BASE_DIR, f"vs_{doc_name}")
                
                # Check if document is already processed
                if doc_name not in self.documents and not os.path.exists(db_path):
                    file_ext = os.path.splitext(file_path)[1].lower()
                    doc_type = {'.pdf': 'pdf', '.txt': 'txt', '.csv': 'csv', '.xlsx': 'xlsx', '.xls': 'xlsx'}.get(file_ext, 'txt')
                    
                    logger.info(f"Auto-loading document: {doc_name}")
                    self.add_document(file_path, doc_type)
                elif doc_name not in self.documents and os.path.exists(db_path):
                    # Document was processed before, just add to registry
                    file_ext = os.path.splitext(file_path)[1].lower()
                    doc_type = {'.pdf': 'pdf', '.txt': 'txt', '.csv': 'csv', '.xlsx': 'xlsx', '.xls': 'xlsx'}.get(file_ext, 'txt')
                    
                    # Count chunks by loading the vector store
                    try:
                        vectordb = FAISS.load_local(
                            db_path,
                            self.embeddings,
                            allow_dangerous_deserialization=True
                        )
                        chunk_count = vectordb.index.ntotal if hasattr(vectordb.index, 'ntotal') else 0
                    except:
                        chunk_count = 0
                    
                    self.documents[doc_name] = {
                        'path': file_path,
                        'type': doc_type,
                        'db_path': db_path,
                        'chunks': chunk_count,
                        'added_at': datetime.fromtimestamp(os.path.getctime(file_path)).isoformat(),
                        'auto_loaded': True
                    }
                    logger.info(f"Registered existing document: {doc_name}")
        
        # Rebuild combined vector store if we have documents
        if self.documents:
            self._rebuild_combined_vectordb()
            logger.info(f"Auto-loaded {len(self.documents)} documents from knowledge base folder")
    
    def scan_for_new_documents(self):
        """Scan for new documents added to the folder"""
        old_count = len(self.documents)
        self.load_existing_documents()
        new_count = len(self.documents)
        return new_count - old_count
    
    def search_knowledge_base(self, question: str, k: int = 3) -> List[Dict]:
        """Search the knowledge base for relevant information"""
        if not self.combined_vectordb:
            return []
        
        try:
            docs = self.combined_vectordb.similarity_search(question, k=k)
            results = []
            for doc in docs:
                results.append({
                    'content': doc.page_content,
                    'metadata': doc.metadata
                })
            return results
        except Exception as e:
            logger.error(f"Error searching knowledge base: {e}")
            return []

class RAGSystem:
    """Standalone RAG System - No Web Interface"""
    
    def __init__(self):
        """Initialize the RAG system"""
        self.kb_manager = KnowledgeBaseManager()
        self.llm = Ollama(model=ANSWER_MODEL)
        print(f"🤖 RAG System initialized with {len(self.kb_manager.documents)} documents")
    
    def ask_question(self, question: str, use_cache: bool = True, k: int = 3) -> Dict:
        """Ask a question to the knowledge base"""
        start_time = time.time()
        
        # Check cache first
        if use_cache:
            cached = self.kb_manager.get_cached_answer(question)
            if cached:
                execution_time = time.time() - start_time
                self.kb_manager.log_operation("cached_answer", f"Question: {question[:50]}...", execution_time)
                print(f"💾 Retrieved cached answer in {execution_time:.2f}s")
                return {
                    "answer": cached['answer'],
                    "cached": True,
                    "timestamp": cached['timestamp'],
                    "execution_time": execution_time,
                    "source_files": cached['source_files']
                }
        
        # Search knowledge base
        search_results = self.kb_manager.search_knowledge_base(question, k=k)
        
        if not search_results:
            return {
                "answer": "No relevant information found in knowledge base",
                "cached": False,
                "execution_time": time.time() - start_time
            }
        
        # Combine context from search results
        context = "\n\n".join([result['content'] for result in search_results])
        source_files = list(set([result['metadata'].get('source', 'unknown') for result in search_results]))
        
        # Generate answer with table and column awareness
        prompt = f"""
You are an intelligent assistant analyzing documents that may contain tables, Excel sheets with multiple columns, and structured data.

IMPORTANT INSTRUCTIONS:
1. If the question asks about specific columns (like 'min', 'max', 'range', 'parameter', etc.), search carefully in ALL provided context
2. If you find table data or Excel columns in the context, extract the relevant values precisely
3. If the data contains numerical values in columns like 'min', 'max', provide the exact numbers
4. If the answer exists in any table, Excel sheet, or structured data, provide it clearly
5. If no relevant information is found, explicitly state: "Information not available in knowledge base"

CONTEXT (includes text, tables, and Excel data):
{context}

QUESTION: {question}

ANSWER (Be specific and include exact values from tables/columns if found):
"""
        
        try:
            print(f"🔍 Generating answer for: {question}")
            answer = self.llm.invoke(prompt)
            
            # Cache the answer
            if use_cache:
                self.kb_manager.cache_answer(question, answer, context, source_files)
            
            execution_time = time.time() - start_time
            self.kb_manager.log_operation("answer_generated", f"Question: {question[:50]}...", execution_time)
            print(f"✅ Answer generated in {execution_time:.2f}s")
            
            return {
                "answer": answer,
                "cached": False,
                "source_files": source_files,
                "execution_time": execution_time
            }
            
        except Exception as e:
            logger.error(f"Error generating answer: {e}")
            return {"error": f"Error generating answer: {str(e)}"}
    
    def add_document(self, file_path: str) -> bool:
        """Add a document to the knowledge base"""
        file_ext = os.path.splitext(file_path)[1].lower()
        doc_type = {'.pdf': 'pdf', '.txt': 'txt', '.csv': 'csv'}.get(file_ext, 'txt')
        return self.kb_manager.add_document(file_path, doc_type)
    
    def scan_for_new_documents(self) -> int:
        """Scan for new documents in the knowledge base folder"""
        return self.kb_manager.scan_for_new_documents()
    
    def get_status(self) -> Dict:
        """Get system status"""
        cache_size = self._get_cache_size()
        return {
            "total_documents": len(self.kb_manager.documents),
            "documents": list(self.kb_manager.documents.keys()),
            "cache_size": cache_size,
            "knowledge_base_dir": KNOWLEDGE_BASE_DIR
        }
    
    def _get_cache_size(self) -> int:
        """Get number of cached answers"""
        conn = sqlite3.connect(CACHE_DB)
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT COUNT(*) FROM cached_answers")
            return cursor.fetchone()[0]
        finally:
            conn.close()
    
    def list_documents(self):
        """List all documents in the knowledge base"""
        if not self.kb_manager.documents:
            print("📁 No documents found in knowledge base")
            return
        
        print(f"📚 Knowledge Base Documents ({len(self.kb_manager.documents)}):")
        print("-" * 50)
        for doc_name, doc_info in self.kb_manager.documents.items():
            status = "🔄 Auto-loaded" if doc_info.get('auto_loaded', False) else "📤 Manually added"
            print(f"{status} | {doc_name} | {doc_info['chunks']} chunks | {doc_info['type'].upper()}")
    
    def fibonacci(self, n: int) -> Dict:
        """Calculate Fibonacci sequence"""
        if n <= 0:
            return {"error": "Number must be positive"}
        
        fib_sequence = [0, 1]
        for i in range(2, n):
            fib_sequence.append(fib_sequence[i-1] + fib_sequence[i-2])
        
        return {
            "n": n,
            "sequence": fib_sequence[:n],
            "nth_number": fib_sequence[n-1] if n <= len(fib_sequence) else None
        }


def demo_usage():
    """Demonstrate how to use the RAG system"""
    print("=" * 60)
    print("🧠 Standalone RAG Knowledge Base System Demo")
    print("=" * 60)
    
    # Initialize the system
    rag = RAGSystem()
    
    # Show system status
    status = rag.get_status()
    print(f"\n📊 System Status:")
    print(f"   Documents: {status['total_documents']}")
    print(f"   Cached answers: {status['cache_size']}")
    print(f"   Knowledge base folder: {status['knowledge_base_dir']}")
    
    # List documents
    print("\n")
    rag.list_documents()
    
    if len(rag.kb_manager.documents) == 0:
        print(f"\n⚠️  No documents found!")
        print(f"📁 Add PDF, TXT, or CSV files to: {KNOWLEDGE_BASE_DIR}/")
        return rag
    
    # Ask some demo questions
    demo_questions = [
        "What is artificial intelligence?",
        "Tell me about machine learning",
        "What is RAG?",
        "How do vector databases work?"
    ]
    
    print(f"\n🤖 Asking demo questions...")
    print("-" * 40)
    
    for question in demo_questions:
        print(f"\n❓ Q: {question}")
        result = rag.ask_question(question)
        
        if 'error' in result:
            print(f"❌ Error: {result['error']}")
        else:
            cached_indicator = "💾 (cached)" if result.get('cached', False) else "🆕 (new)"
            print(f"🤖 A: {result['answer'][:200]}...")
            print(f"⏱️  Time: {result['execution_time']:.2f}s {cached_indicator}")
            if result.get('source_files'):
                print(f"📄 Sources: {', '.join(result['source_files'])}")
    
    # Fibonacci demo
    print(f"\n🧮 Fibonacci Demo:")
    fib_result = rag.fibonacci(10)
    print(f"First 10 Fibonacci numbers: {fib_result['sequence']}")
    
    return rag


def interactive_mode(rag_system):
    """Interactive question-answer mode"""
    print("\n" + "=" * 60)
    print("💬 Interactive Mode - Type your questions!")
    print("Commands: 'quit', 'exit', 'status', 'docs', 'scan'")
    print("=" * 60)
    
    while True:
        try:
            question = input("\n❓ Your question: ").strip()
            
            if question.lower() in ['quit', 'exit', 'bye']:
                print("👋 Goodbye!")
                break
            elif question.lower() == 'status':
                status = rag_system.get_status()
                print(f"📊 Documents: {status['total_documents']}, Cache: {status['cache_size']}")
                continue
            elif question.lower() == 'docs':
                rag_system.list_documents()
                continue
            elif question.lower() == 'scan':
                new_docs = rag_system.scan_for_new_documents()
                print(f"🔍 Found {new_docs} new documents")
                continue
            elif question == '':
                continue
            
            # Ask the question
            result = rag_system.ask_question(question)
            
            if 'error' in result:
                print(f"❌ Error: {result['error']}")
            else:
                cached_indicator = "💾 (cached)" if result.get('cached', False) else "🆕 (new)"
                print(f"\n🤖 Answer: {result['answer']}")
                print(f"⏱️  Time: {result['execution_time']:.2f}s {cached_indicator}")
                if result.get('source_files'):
                    print(f"📄 Sources: {', '.join(result['source_files'])}")
        
        except KeyboardInterrupt:
            print("\n👋 Goodbye!")
            break
        except Exception as e:
            print(f"❌ Error: {e}")


if __name__ == "__main__":
    """Main entry point"""
    import sys
    
    # Create knowledge base folder if it doesn't exist
    os.makedirs(KNOWLEDGE_BASE_DIR, exist_ok=True)
    
    if len(sys.argv) > 1:
        # Command line usage
        command = sys.argv[1].lower()
        
        if command == "demo":
            rag = demo_usage()
            if len(sys.argv) > 2 and sys.argv[2].lower() == "interactive":
                interactive_mode(rag)
        elif command == "ask" and len(sys.argv) > 2:
            # Direct question from command line
            rag = RAGSystem()
            question = " ".join(sys.argv[2:])
            result = rag.ask_question(question)
            if 'error' in result:
                print(f"Error: {result['error']}")
            else:
                print(result['answer'])
        elif command == "interactive":
            rag = RAGSystem()
            interactive_mode(rag)
        else:
            print("Usage:")
            print("  python rag_pdf.py demo                    # Run demo")
            print("  python rag_pdf.py demo interactive        # Demo + interactive mode")
            print("  python rag_pdf.py interactive             # Interactive mode only")
            print("  python rag_pdf.py ask \"your question\"     # Ask single question")
    else:
        # Default: run demo
        rag = demo_usage()
        interactive_mode(rag)